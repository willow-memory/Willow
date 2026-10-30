"""sheets: each sheet's cells by address, shown as their number formats show
them, formulas after their cached results, notes after the rows; the shared
reader contract; and a refusal, never a quiet skip, without openpyxl."""

from __future__ import annotations

import datetime as dt
import importlib.util
import json
import subprocess
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
KERNEL = HERE.parents[2] / "kernel.py"
SCRIPT = HERE / "sheets.py"
spec = importlib.util.spec_from_file_location("sheets", SCRIPT)
sh = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sh)
NOW = ["--now", "2026-10-29T09:00:00Z"]


def cached(path: Path, cell: str, value: str) -> None:
    """Give `cell`'s formula the result a spreadsheet program would have
    saved beside it; openpyxl saves none."""
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    name = "xl/worksheets/sheet1.xml"
    xml = parts[name].decode()
    at = xml.index(f'<c r="{cell}"')
    end = xml.index("</c>", at)
    cut = xml[at:end]
    done = cut.replace("<v />", f"<v>{value}</v>").replace("<v></v>", f"<v>{value}</v>")
    if done == cut:
        done = cut + f"<v>{value}</v>"
    parts[name] = (xml[:at] + done + xml[end:]).encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)


def budget(path: Path) -> None:
    import openpyxl
    from openpyxl.comments import Comment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Cost", "Due"])
    ws.append(["Rent", 1200, dt.date(2026, 3, 1)])
    ws.append(["Gas", 250.5, None])
    ws.append(["Total", "=SUM(B2:B3)", None])
    for c in ("B2", "B3", "B4"):
        ws[c].number_format = '"£"#,##0.00'
    ws["C2"].number_format = "d/m/yyyy"
    ws["B4"].comment = Comment("check the\ndeposit", "Ana")
    rates = wb.create_sheet("Rates")
    rates["A1"], rates["B1"] = "Interest", 0.0425
    rates["B1"].number_format = "0.00%"
    rates.sheet_state = "hidden"
    wb.save(path)
    cached(path, "B4", "1450.5")


def run(*argv: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, *argv], capture_output=True, text=True)


def folder(tmp_path: Path) -> Path:
    src = tmp_path / "money"
    (src / "2026").mkdir(parents=True)
    budget(src / "2026" / "budget.xlsx")
    (src / "broken.xlsx").write_bytes(b"PK not really a zip")
    (src / "2026" / "~$budget.xlsx").write_bytes(b"Excel's lock file")
    (src / ".hidden").mkdir()
    budget(src / ".hidden" / "x.xlsx")
    (src / "notes.txt").write_text("not a sheet\n")
    return src


def mirror(src: Path, dst: Path) -> subprocess.CompletedProcess:
    return run(str(SCRIPT), "mirror", "--from", str(src), "--to", str(dst))


def test_each_sheet_is_a_passage_of_cells_as_a_person_sees_them(tmp_path):
    src, dst = folder(tmp_path), tmp_path / "mirror"
    got = mirror(src, dst)
    assert got.returncode == 0, got.stderr
    assert got.stdout.startswith("wrote 1 files to")
    assert got.stdout.strip().endswith("skipped 1")
    assert (dst / "2026" / "budget.xlsx.txt").read_text() == (
        "== sheet Budget\n"
        "A1 Item | B1 Cost | C1 Due\n"
        "A2 Rent | B2 £1,200.00 | C2 2026-03-01\n"
        "A3 Gas | B3 £250.50\n"
        "A4 Total | B4 £1,450.50 =SUM(B2:B3)\n"
        "note B4 (Ana): check the deposit\n"
        "\n"
        "== sheet Rates (hidden)\n"
        "A1 Interest | B1 4.25%\n"
    )
    assert sorted(p.name for p in dst.rglob("*") if p.is_file()) == ["budget.xlsx.txt"]


def test_a_formula_with_no_saved_result_shows_alone(tmp_path):
    import openpyxl

    src = tmp_path / "s"
    src.mkdir()
    wb = openpyxl.Workbook()
    wb.active["A1"], wb.active["A2"] = 2, "=A1*3"
    wb.save(src / "calc.xlsx")
    assert mirror(src, tmp_path / "m").returncode == 0
    assert (tmp_path / "m" / "calc.xlsx.txt").read_text() == (
        "== sheet Sheet\nA1 2\nA2 =A1*3\n"
    )


def test_number_formats_as_a_spreadsheet_shows_them():
    cases = [
        (1200, '"£"#,##0.00', "£1,200.00"),
        (250.5, "[$€-407]#,##0.00;[Red]-[$€-407]#,##0.00", "€250.50"),
        (0.125, "0.0%", "12.5%"),
        (-1234.5, "#,##0", "-1,235"),  # half away from zero, as Excel rounds
        (2.5, "0", "3"),
        (12345.678, "0.00E+00", "1.23E+04"),
        (5, '0" kg"', "5 kg"),
        (1200, "_($* #,##0.00_)", "$1,200.00"),
        (0.5, "# ?/?", "0.5"),  # a fraction it doesn't draw: as General
        (-0.001, "0.00", "0.00"),  # no minus on a zero
    ]
    for v, fmt, want in cases:
        assert sh.number(v, fmt) == want, (v, fmt)
    assert sh.shown(3.0, "General") == "3" and sh.shown(0.1, "General") == "0.1"
    assert sh.shown(True, "General") == "TRUE"
    assert sh.shown(dt.datetime(2026, 1, 2, 3, 4, 5), "") == "2026-01-02 03:04:05"


def test_the_kernel_finds_a_cell_by_its_value_and_its_note(tmp_path):
    src, dst = folder(tmp_path), tmp_path / "mirror"
    mirror(src, dst)
    b = tmp_path / "box"
    run(str(KERNEL), "init", "--box", str(b), "--pile", str(dst),
        "--key", str(tmp_path / "k"), *NOW)  # fmt: skip
    got = json.loads(
        run(str(KERNEL), "bite", "--box", str(b), *NOW, "total deposit").stdout
    )
    assert got["state"] == "populated"
    assert got["served"][0]["where"] == "2026/budget.xlsx.txt:1"
    assert "£1,450.50" in got["served"][0]["text"]


def test_the_same_file_gives_the_same_bytes_and_originals_are_untouched(tmp_path):
    src = folder(tmp_path)
    before = {p: p.read_bytes() for p in src.rglob("*") if p.is_file()}
    for name in ("one", "two"):
        mirror(src, tmp_path / name)
    a = (tmp_path / "one" / "2026" / "budget.xlsx.txt").read_bytes()
    assert a == (tmp_path / "two" / "2026" / "budget.xlsx.txt").read_bytes()
    assert {p: p.read_bytes() for p in src.rglob("*") if p.is_file()} == before


def test_refusals_and_nothing_readable(tmp_path):
    src = folder(tmp_path)
    assert mirror(src, src / "m").returncode == 2
    assert mirror(src, tmp_path).returncode == 2
    assert run(str(SCRIPT), "mirror", "--from", str(src)).returncode == 2
    empty = tmp_path / "empty"
    empty.mkdir()
    (empty / "bad.xlsx").write_bytes(b"nope")
    assert mirror(empty, tmp_path / "m2").returncode == 42


def test_without_openpyxl_it_refuses_and_names_the_install(tmp_path):
    blocker = tmp_path / "block"
    blocker.mkdir()
    (blocker / "openpyxl.py").write_text("raise ImportError('hidden')\n")
    src = folder(tmp_path)
    script = (
        "import runpy, sys\n"
        f"sys.path.insert(0, {str(blocker)!r})\n"
        f"sys.argv = [{str(SCRIPT)!r}] + sys.argv[1:]\n"
        f"runpy.run_path({str(SCRIPT)!r}, run_name='__main__')\n"
    )
    got = run("-c", script, "mirror", "--from", str(src), "--to", str(tmp_path / "m"))
    assert got.returncode == 2 and "requirements.txt" in got.stderr
